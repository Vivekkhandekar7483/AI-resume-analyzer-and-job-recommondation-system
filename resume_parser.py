import io
from pypdf import PdfReader
from docx import Document


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract raw text from PDF or DOCX file bytes entirely in memory."""
    if not file_bytes:
        raise ValueError("File is empty.")

    name = filename.lower()
    if name.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            return "\n".join(pages_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {e}")

    if name.endswith(".docx"):
        try:
            doc = Document(io.BytesIO(file_bytes))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        parts.append(row_text)
            return "\n".join(parts)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {e}")

    raise ValueError("Only PDF (.pdf) and DOCX (.docx) files are supported.")

