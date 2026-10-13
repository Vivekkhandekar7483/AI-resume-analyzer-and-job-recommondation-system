import os
from docx import Document

def build_resume(filename, title, summary, skills, experience, education):
    doc = Document()
    doc.add_heading(title, level=0)
    
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(summary)
    
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(skills)
    
    doc.add_heading("Projects & Experience", level=1)
    for exp in experience:
        doc.add_heading(exp["title"], level=2)
        doc.add_paragraph(exp["desc"])
        
    doc.add_heading("Education", level=1)
    doc.add_paragraph(education)
    
    doc.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    out_dir = os.path.dirname(__file__)
    
    # 1. Data Analyst Sample
    build_resume(
        os.path.join(out_dir, "Resume_Data_Analyst.docx"),
        "Alex Rivera - Data Analyst",
        "Detail-oriented Data Analyst with 2+ years experience in SQL queries, Python data wrangling, and Power BI dashboarding.",
        "Programming: Python, SQL\nData: Pandas, NumPy, Data Visualization, Statistics, Excel\nTools: Power BI, Git, Tableau",
        [
            {"title": "Sales Analytics Dashboard", "desc": "Built Power BI dashboards analyzing quarterly sales metrics using SQL joins and Pandas wrangling."},
            {"title": "Customer Churn Analysis", "desc": "Performed statistical analysis and data visualization on customer churn dataset using Python and Seaborn."}
        ],
        "B.S. in Computer Science - State University (2024)"
    )

    # 2. ML Engineer Sample
    build_resume(
        os.path.join(out_dir, "Resume_ML_Engineer.docx"),
        "Jordan Lee - Machine Learning Engineer",
        "Passionate Machine Learning Engineer specializing in scikit-learn models, FastAPI microservices, and Docker containerization.",
        "ML & AI: Machine Learning, scikit-learn, Deep Learning, TensorFlow, PyTorch\nProgramming: Python, C++, SQL\nDevOps & Web: FastAPI, Docker, MLflow, Git, REST API",
        [
            {"title": "Predictive Maintenance Pipeline", "desc": "Trained predictive ML models with scikit-learn and deployed REST APIs containerized using Docker and FastAPI."},
            {"title": "MLOps Experimentation", "desc": "Tracked model hyperparameter runs using MLflow and deployed pipeline on Linux server."}
        ],
        "B.Tech in Artificial Intelligence - Tech University (2023)"
    )

    # 3. NLP Engineer Sample
    build_resume(
        os.path.join(out_dir, "Resume_NLP_Engineer.docx"),
        "Sam Taylor - NLP Engineer",
        "NLP Engineer experienced in Transformer models, Hugging Face pipelines, spaCy text processing, and Deep Learning.",
        "NLP & ML: NLP, Transformers, Hugging Face, spaCy, Deep Learning, PyTorch, LLM, RAG\nProgramming: Python, SQL\nTools: Git, Docker, REST API",
        [
            {"title": "Text Summarization Engine", "desc": "Fine-tuned BERT and Transformer models using Hugging Face for automated document summarization."},
            {"title": "RAG Question Answering System", "desc": "Built retrieval augmented generation pipeline using LangChain, spaCy, and vector embeddings."}
        ],
        "M.S. in Computational Linguistics - Innovation Institute (2024)"
    )
